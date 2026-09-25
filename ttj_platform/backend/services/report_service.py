import io
import csv
import datetime
from typing import Dict, Any, List, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from backend.models.database import get_db, dicts_from_rows
from backend.auth.security import mask_jshshir

REPORT_TITLES = {
    "all_students": "1. Barcha talabalar ro'yxati",
    "rooms_list": "2. TTJ xonalar va sig'im ro'yxati",
    "vacant_beds": "3. Bo'sh joylar ro'yxati",
    "occupied_beds": "4. Band joylar ro'yxati",
    "by_faculty": "5. Fakultetlar bo'yicha hisobot",
    "by_course": "6. Kurslar bo'yicha hisobot",
    "by_privilege": "7. Imtiyozli talabalar hisoboti",
    "payments": "8. To'lovlar jurnali",
    "debtors": "9. Qarzdor talabalar ro'yxati",
    "moved_students": "10. Xonasi almashtirilgan talabalar",
    "vacated_students": "11. TTJdan chiqarilgan talabalar",
    "change_history": "12. Tizimdagi o'zgarishlar tarixi"
}

def get_report_data(report_type: str) -> Tuple[str, List[str], List[List[Any]]]:
    title = REPORT_TITLES.get(report_type, "TTJ Hisoboti")
    headers = []
    rows = []
    
    with get_db() as conn:
        if report_type == "all_students":
            headers = ["№", "F.I.O.", "JShShIR", "Fakultet", "Kurs", "Bino", "Xona", "O'rin", "To'lov Holati", "Qarzdorlik (so'm)"]
            query = """
                SELECT s.full_name, s.jshshir, s.faculty, s.course, b.name as building_name,
                       r.room_number, rb.bed_number, s.payment_status, s.remaining_fee
                FROM students s
                LEFT JOIN rooms r ON s.current_room_id = r.id
                LEFT JOIN room_beds rb ON s.current_bed_id = rb.id
                LEFT JOIN buildings b ON r.building_id = b.id
                WHERE s.status = 'active' AND s.deleted_at IS NULL
                ORDER BY s.full_name ASC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                p_status = "To'liq" if r['payment_status'] == 'paid' else ("Qisman" if r['payment_status'] == 'partial' else "To'lanmagan")
                rows.append([idx, r['full_name'], mask_jshshir(r['jshshir']), r['faculty'], f"{r['course']}-kurs",
                             r['building_name'] or "-", r['room_number'] or "-", r['bed_number'] or "-",
                             p_status, f"{r['remaining_fee']:,.0f}"])
                             
        elif report_type == "rooms_list":
            headers = ["№", "Bino", "Blok", "Qavat", "Xona", "Turi", "Sig'im", "Band", "Bo'sh", "Bandlik %"]
            query = """
                SELECT b.name as building_name, blk.name as block_name, f.floor_number, r.room_number, r.room_type, r.capacity,
                       (SELECT COUNT(*) FROM room_beds rb WHERE rb.room_id = r.id AND rb.status = 'occupied') as occ,
                       (SELECT COUNT(*) FROM room_beds rb WHERE rb.room_id = r.id AND rb.status = 'vacant') as vac
                FROM rooms r
                JOIN buildings b ON r.building_id = b.id
                JOIN blocks blk ON r.block_id = blk.id
                JOIN floors f ON r.floor_id = f.id
                WHERE r.deleted_at IS NULL
                ORDER BY b.number, blk.block_number, f.floor_number, r.room_number
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                pct = round((r['occ'] / r['capacity']) * 100, 1) if r['capacity'] > 0 else 0
                rows.append([idx, r['building_name'], r['block_name'], f"{r['floor_number']}-qavat", r['room_number'],
                             r['room_type'].capitalize(), r['capacity'], r['occ'], r['vac'], f"{pct}%"])
                             
        elif report_type == "vacant_beds":
            headers = ["№", "Bino", "Blok", "Qavat", "Xona", "Bo'sh o'rin raqami"]
            query = """
                SELECT b.name as b_name, blk.name as blk_name, f.floor_number, r.room_number, rb.bed_number
                FROM room_beds rb
                JOIN rooms r ON rb.room_id = r.id
                JOIN buildings b ON r.building_id = b.id
                JOIN blocks blk ON r.block_id = blk.id
                JOIN floors f ON r.floor_id = f.id
                WHERE rb.status = 'vacant' AND r.deleted_at IS NULL
                ORDER BY b.number, blk.block_number, f.floor_number, r.room_number, rb.bed_number
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['b_name'], r['blk_name'], f"{r['floor_number']}-qavat", r['room_number'], f"{r['bed_number']}-o'rin"])
                
        elif report_type == "occupied_beds":
            headers = ["№", "Bino", "Xona", "O'rin", "Talaba F.I.O.", "Fakultet", "Kurs", "Telefon"]
            query = """
                SELECT b.name as b_name, r.room_number, rb.bed_number, s.full_name, s.faculty, s.course, s.phone
                FROM room_beds rb
                JOIN rooms r ON rb.room_id = r.id
                JOIN buildings b ON r.building_id = b.id
                JOIN students s ON rb.current_student_id = s.id
                WHERE rb.status = 'occupied' AND s.status = 'active'
                ORDER BY b.number, r.room_number, rb.bed_number
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['b_name'], r['room_number'], f"{r['bed_number']}-o'rin", r['full_name'], r['faculty'], f"{r['course']}-kurs", r['phone']])
                
        elif report_type == "by_faculty":
            headers = ["№", "Fakultet", "Yashovchi talabalar soni", "To'langan summa (so'm)", "Qarzdorlik (so'm)"]
            query = """
                SELECT faculty, COUNT(*) as count, SUM(paid_fee) as total_paid, SUM(remaining_fee) as total_rem
                FROM students WHERE status = 'active' GROUP BY faculty ORDER BY count DESC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['faculty'], r['count'], f"{r['total_paid'] or 0:,.0f}", f"{r['total_rem'] or 0:,.0f}"])
                
        elif report_type == "by_course":
            headers = ["№", "Kurs", "Talabalar soni", "To'langan summa (so'm)", "Qarzdorlik (so'm)"]
            query = """
                SELECT course, COUNT(*) as count, SUM(paid_fee) as total_paid, SUM(remaining_fee) as total_rem
                FROM students WHERE status = 'active' GROUP BY course ORDER BY course ASC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, f"{r['course']}-kurs", r['count'], f"{r['total_paid'] or 0:,.0f}", f"{r['total_rem'] or 0:,.0f}"])
                
        elif report_type == "by_privilege":
            headers = ["№", "Imtiyoz turi", "Talabalar soni", "Talabalar ro'yxati (F.I.O. namuna)"]
            query = """
                SELECT p.name as priv_name, COUNT(s.id) as count, GROUP_CONCAT(s.full_name, ', ') as names
                FROM student_privileges p
                LEFT JOIN students s ON s.privilege_id = p.id AND s.status = 'active'
                GROUP BY p.id ORDER BY count DESC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                sample_names = (r['names'] or "")[:60] + ("..." if r['names'] and len(r['names']) > 60 else "")
                rows.append([idx, r['priv_name'], r['count'], sample_names or "Talaba yo'q"])
                
        elif report_type == "payments":
            headers = ["№", "Sana", "Talaba F.I.O.", "Summa (so'm)", "To'lov turi", "Kvitansiya №", "Qabul qildi"]
            query = """
                SELECT p.*, s.full_name FROM payments p
                JOIN students s ON p.student_id = s.id
                ORDER BY p.payment_date DESC, p.id DESC LIMIT 100
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['payment_date'], r['full_name'], f"{r['amount']:,.0f}", r['payment_method'], r['receipt_number'] or "-", r['created_by'] or "-"])
                
        elif report_type == "debtors":
            headers = ["№", "Talaba F.I.O.", "Telefon", "Fakultet", "Kurs", "Xona", "Umumiy to'lov", "To'langan", "Qarzdorlik"]
            query = """
                SELECT s.*, r.room_number, b.name as building_name
                FROM students s
                LEFT JOIN rooms r ON s.current_room_id = r.id
                LEFT JOIN buildings b ON r.building_id = b.id
                WHERE s.remaining_fee > 0 AND s.status = 'active'
                ORDER BY s.remaining_fee DESC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['full_name'], r['phone'], r['faculty'], f"{r['course']}-kurs",
                             f"{r['building_name']} {r['room_number']}", f"{r['total_fee']:,.0f}", f"{r['paid_fee']:,.0f}", f"{r['remaining_fee']:,.0f}"])
                             
        elif report_type == "moved_students":
            headers = ["№", "Talaba F.I.O.", "Sana", "Xonadan -> Xonaga", "Sabab", "Mas'ul xodim"]
            query = """
                SELECT sh.*, s.full_name
                FROM student_history sh
                JOIN students s ON sh.student_id = s.id
                WHERE sh.action_type = 'move'
                ORDER BY sh.id DESC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['full_name'], r['created_at'][:16], f"{r['old_value']} -> {r['new_value']}", r['reason'] or "-", r['user_name'] or "-"])
                
        elif report_type == "vacated_students":
            headers = ["№", "Talaba F.I.O.", "JShShIR", "Fakultet", "Chiqarilgan sana", "Chiqarish sababi", "Holat"]
            query = """
                SELECT s.full_name, s.jshshir, s.faculty, s.check_out_date, s.vacate_reason, s.status
                FROM students s
                WHERE s.status IN ('vacated', 'graduated', 'academic_leave')
                ORDER BY s.updated_at DESC
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                status_uz = "TTJdan chiqarilgan" if r['status'] == 'vacated' else ("O'qishni tugatgan" if r['status'] == 'graduated' else "Akademik ta'tilda")
                rows.append([idx, r['full_name'], mask_jshshir(r['jshshir']), r['faculty'], r['check_out_date'] or "-", r['vacate_reason'] or "-", status_uz])
                
        elif report_type == "change_history":
            headers = ["№", "Sana/Vaqt", "Amal turi", "Ob'yekt", "Foydalanuvchi", "Izoh/Sabab"]
            query = """
                SELECT al.* FROM audit_logs al ORDER BY al.id DESC LIMIT 100
            """
            for idx, r in enumerate(conn.execute(query).fetchall(), 1):
                rows.append([idx, r['created_at'][:16], r['action'], f"{r['entity_type']} #{r['entity_id'] or ''}", r['user_name'] or "Tizim", r['reason'] or "-"])
                
    return title, headers, rows

def export_report_csv(report_type: str) -> str:
    title, headers, rows = get_report_data(report_type)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([title])
    writer.writerow([f"Shakllantirildi: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"])
    writer.writerow([])
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return output.getvalue()

def export_report_excel(report_type: str) -> bytes:
    title, headers, rows = get_report_data(report_type)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    
    # Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title.upper()
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Date row
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    date_cell = ws.cell(row=2, column=1)
    date_cell.value = f"Shakllantirilgan sana: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"
    date_cell.font = Font(name="Calibri", size=10, italic=True, color="555555")
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Headers
    h_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    for c_idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=c_idx)
        c.value = h
        c.fill = h_fill
        c.font = h_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        
    # Rows
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for r_idx, row_data in enumerate(rows, 5):
        for c_idx, val in enumerate(row_data, 1):
            c = ws.cell(row=r_idx, column=c_idx)
            c.value = val
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = zebra_fill
            if c_idx == 1:
                c.alignment = Alignment(horizontal="center")
                
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def export_report_pdf(report_type: str) -> bytes:
    title, headers, rows = get_report_data(report_type)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1, # Center
        spaceAfter=5
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor('#666666'),
        alignment=1,
        spaceAfter=15
    )
    
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=8,
        leading=10
    )
    header_cell_style = ParagraphStyle(
        'HeaderCellText',
        parent=styles['Normal'],
        fontSize=8,
        fontName='Helvetica-Bold',
        textColor=colors.white,
        leading=10,
        alignment=1
    )
    
    elements = []
    elements.append(Paragraph("TALABALAR TURAR JOYI (TTJ) BOSHQARUV TIZIMI", title_style))
    elements.append(Paragraph(f"{title} &bull; Sana: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}", sub_style))
    
    # Table data
    table_data = []
    # Header row
    h_row = [Paragraph(str(h), header_cell_style) for h in headers]
    table_data.append(h_row)
    
    for row in rows[:150]: # Limit for PDF clean rendering
        t_row = [Paragraph(str(val), cell_style) for val in row]
        table_data.append(t_row)
        
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    elements.append(t)
    doc.build(elements)
    return out.getvalue()
