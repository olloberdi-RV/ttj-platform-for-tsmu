import os
import io
import csv
import json
import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import docx
from fastapi import HTTPException, status
from backend.models.database import get_db, dicts_from_rows, dict_from_row
from backend.auth.security import validate_jshshir, mask_jshshir
from backend.services.audit_service import log_audit

def parse_docx_bytes(content: bytes) -> List[Dict[str, Any]]:
    doc = docx.Document(io.BytesIO(content))
    rows_data = []
    
    # Check tables
    if not doc.tables:
        return []
        
    table = doc.tables[0]
    headers = []
    for i, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        if i == 0:
            headers = [h.lower() for h in cells]
            continue
        if not any(cells):
            continue
        row_dict = {}
        for idx, val in enumerate(cells):
            h = headers[idx] if idx < len(headers) else f"col_{idx}"
            row_dict[h] = val
        rows_data.append(row_dict)
    return rows_data

def parse_xlsx_bytes(content: bytes) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows_data = []
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip().lower() if h is not None else f"col_{idx}" for idx, h in enumerate(row)]
            continue
        if not any(row):
            continue
        row_dict = {}
        for idx, val in enumerate(row):
            h = headers[idx] if idx < len(headers) else f"col_{idx}"
            row_dict[h] = str(val).strip() if val is not None else ""
        rows_data.append(row_dict)
    return rows_data

def parse_csv_bytes(content: bytes) -> List[Dict[str, Any]]:
    decoded = content.decode('utf-8-sig', errors='replace')
    reader = csv.reader(io.StringIO(decoded))
    rows_data = []
    headers = []
    for i, row in enumerate(reader):
        if i == 0:
            headers = [h.strip().lower() for h in row]
            continue
        if not any(row):
            continue
        row_dict = {}
        for idx, val in enumerate(row):
            h = headers[idx] if idx < len(headers) else f"col_{idx}"
            row_dict[h] = val.strip()
        rows_data.append(row_dict)
    return rows_data

def normalize_row_keys(raw_row: Dict[str, Any]) -> Dict[str, Any]:
    norm = {}
    for k, v in raw_row.items():
        key = k.lower().replace(" ", "").replace("_", "").replace(".", "").replace("‘", "").replace("'", "")
        if "fio" in key or "ism" in key or "talaba" in key:
            norm["full_name"] = v
        elif "jshshir" in key or "pinfl" in key or "pasport" in key:
            norm["jshshir"] = v
        elif "fakultet" in key:
            norm["faculty"] = v
        elif "kurs" in key:
            norm["course"] = v
        elif "xona" in key:
            norm["room_number"] = v
        elif "sport" in key:
            norm["sport"] = v
        elif "imtiyoz" in key:
            norm["privilege"] = v
        elif "umumiy" in key or "tolov" in key or "summa" in key:
            if "tolangan" not in key:
                norm["total_fee"] = v
        elif "tolangan" in key:
            norm["paid_fee"] = v
        elif "telefon" in key or "tel" in key:
            norm["phone"] = v
        elif "jinsi" in key:
            norm["gender"] = v
            
    # defaults
    norm.setdefault("full_name", raw_row.get("f.i.o.", raw_row.get("fio", "")))
    norm.setdefault("jshshir", raw_row.get("jshshir", ""))
    norm.setdefault("faculty", raw_row.get("fakultet", "Davolash"))
    norm.setdefault("course", raw_row.get("kurs", "1"))
    norm.setdefault("room_number", raw_row.get("xona", ""))
    norm.setdefault("sport", raw_row.get("sport", "Futbol"))
    norm.setdefault("privilege", raw_row.get("imtiyoz", "Oddiy / Imtiyozsiz"))
    norm.setdefault("total_fee", raw_row.get("umumiy to'lov", "2400000"))
    norm.setdefault("paid_fee", raw_row.get("to'langan", "0"))
    norm.setdefault("phone", raw_row.get("telefon", "+998901234567"))
    norm.setdefault("gender", raw_row.get("jinsi", "Erkak"))
    return norm

def validate_import_data(raw_rows: List[Dict[str, Any]], file_name: str) -> Dict[str, Any]:
    with get_db() as conn:
        # Load all active students JShShIRs
        active_jshshirs = {r['jshshir'] for r in conn.execute("SELECT jshshir FROM students WHERE status = 'active' AND deleted_at IS NULL").fetchall()}
        
        # Load all rooms and their free beds count
        rooms_db = conn.execute("""
            SELECT r.id, r.room_number, r.capacity, b.name as building_name,
                   (SELECT COUNT(*) FROM room_beds rb WHERE rb.room_id = r.id AND rb.status = 'vacant') as vacant_count
            FROM rooms r
            JOIN buildings b ON r.building_id = b.id
            WHERE r.deleted_at IS NULL
        """).fetchall()
        
        rooms_map = {}
        for r in rooms_db:
            # Map room_number to room data
            rooms_map[str(r['room_number']).strip().lower()] = {
                "id": r['id'],
                "room_number": r['room_number'],
                "capacity": r['capacity'],
                "vacant_count": r['vacant_count'],
                "allocated_in_batch": 0,
                "building_name": r['building_name']
            }
            
        # Load privileges
        privileges_list = dicts_from_rows(conn.execute("SELECT id, name FROM student_privileges").fetchall())
        priv_map = {p['name'].lower(): p['id'] for p in privileges_list}
        
        # Load sports
        sports_list = dicts_from_rows(conn.execute("SELECT id, name FROM sports").fetchall())
        sport_map = {s['name'].lower(): s['id'] for s in sports_list}

    seen_batch_jshshir = set()
    validated_rows = []
    error_count = 0
    valid_count = 0
    
    for idx, raw in enumerate(raw_rows, start=1):
        item = normalize_row_keys(raw)
        row_errors = []
        
        # 1. Majburiy maydonlar
        if not item.get("full_name") or len(item["full_name"].strip()) < 3:
            row_errors.append("F.I.O. to'liq kiritilishi shart")
            
        jshshir = str(item.get("jshshir", "")).strip().replace(" ", "")
        if not jshshir:
            row_errors.append("JShShIR ko'rsatilmagan")
        elif not validate_jshshir(jshshir):
            row_errors.append("Noto'g'ri JShShIR: aynan 14 ta raqam bo'lishi kerak")
        elif jshshir in seen_batch_jshshir:
            row_errors.append(f"Takroriy JShShIR: ushbu faylning o'zida bir necha bor uchradi ({jshshir})")
        elif jshshir in active_jshshirs:
            row_errors.append(f"Takroriy JShShIR: Ushbu JShShIR bo'yicha faol talaba bazada allaqachon mavjud ({jshshir})")
        else:
            seen_batch_jshshir.add(jshshir)
            
        # Kurs validation
        try:
            course = int(item.get("course", 1))
            if course < 1 or course > 6:
                row_errors.append("Noto'g'ri kurs: kurs 1 dan 6 gacha bo'lishi kerak")
        except (ValueError, TypeError):
            row_errors.append("Noto'g'ri kurs raqami")
            course = 1
            
        # Room validation
        room_num = str(item.get("room_number", "")).strip().lower()
        if not room_num:
            row_errors.append("Xona raqami ko'rsatilmagan")
            room_info = None
        elif room_num not in rooms_map:
            row_errors.append(f"Mavjud bo'lmagan xona: '{item.get('room_number')}' bazada topilmadi")
            room_info = None
        else:
            room_info = rooms_map[room_num]
            if (room_info["vacant_count"] - room_info["allocated_in_batch"]) <= 0:
                row_errors.append(f"Sig'imi to'lgan xona: {room_info['room_number']}-xonada bo'sh o'rin qolmadi (Jami sig'im: {room_info['capacity']})")
            else:
                room_info["allocated_in_batch"] += 1
                
        # Moliyaviy summa validation
        try:
            total_fee = float(str(item.get("total_fee", 2400000)).replace(" ", "").replace(",", "."))
            if total_fee < 0:
                row_errors.append("Umumiy to'lov manfiy bo'lishi mumkin emas")
        except Exception:
            row_errors.append("Umumiy to'lov formati noto'g'ri")
            total_fee = 2400000.0
            
        try:
            paid_fee = float(str(item.get("paid_fee", 0)).replace(" ", "").replace(",", "."))
            if paid_fee < 0:
                row_errors.append("To'langan summa manfiy bo'lishi mumkin emas")
        except Exception:
            row_errors.append("To'langan summa formati noto'g'ri")
            paid_fee = 0.0
            
        is_valid = len(row_errors) == 0
        if is_valid:
            valid_count += 1
        else:
            error_count += 1
            
        validated_rows.append({
            "row_num": idx,
            "full_name": item.get("full_name", "").strip(),
            "jshshir": jshshir,
            "masked_jshshir": mask_jshshir(jshshir),
            "faculty": item.get("faculty", "Davolash"),
            "course": course,
            "room_number": item.get("room_number", ""),
            "room_id": room_info["id"] if room_info else None,
            "sport": item.get("sport", "Futbol"),
            "privilege": item.get("privilege", "Oddiy / Imtiyozsiz"),
            "total_fee": total_fee,
            "paid_fee": paid_fee,
            "phone": item.get("phone", "+998901234567"),
            "gender": item.get("gender", "Erkak"),
            "is_valid": is_valid,
            "errors": row_errors
        })
        
    return {
        "file_name": file_name,
        "total_rows": len(raw_rows),
        "valid_count": valid_count,
        "error_count": error_count,
        "rows": validated_rows
    }

def confirm_import(validated_rows: List[Dict[str, Any]], file_name: str, file_type: str, current_user: Dict[str, Any]) -> Dict[str, Any]:
    with get_db() as conn:
        # Create import session entry
        cursor = conn.execute("""
            INSERT INTO imports (file_name, file_type, user_id, user_name, total_rows, success_count, error_count, status)
            VALUES (?, ?, ?, ?, ?, 0, 0, 'completed')
        """, (file_name, file_type, current_user.get("id"), current_user.get("full_name"), len(validated_rows)))
        import_id = cursor.lastrowid
        
        # Load sport and privilege mappings
        sports_map = {r['name'].lower(): r['id'] for r in dicts_from_rows(conn.execute("SELECT id, name FROM sports").fetchall())}
        priv_map = {r['name'].lower(): r['id'] for r in dicts_from_rows(conn.execute("SELECT id, name FROM student_privileges").fetchall())}
        
        success_count = 0
        error_count = 0
        now_date = datetime.date.today().isoformat()
        
        for r in validated_rows:
            if not r.get("is_valid"):
                error_count += 1
                conn.execute("""
                    INSERT INTO import_errors (import_id, row_number, student_name, jshshir, error_message, raw_data)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    import_id,
                    r.get("row_num", 0),
                    r.get("full_name", ""),
                    r.get("jshshir", ""),
                    "; ".join(r.get("errors", [])),
                    json.dumps(r, ensure_ascii=False)
                ))
                continue
                
            room_id = r.get("room_id")
            # Find first vacant bed in that room
            bed = conn.execute("""
                SELECT id, bed_number FROM room_beds WHERE room_id = ? AND status = 'vacant' ORDER BY bed_number ASC LIMIT 1
            """, (room_id,)).fetchone()
            
            if not bed:
                error_count += 1
                msg = f"Import jarayonida xonada bo'sh joy qolmadi: Xona {r.get('room_number')}"
                conn.execute("""
                    INSERT INTO import_errors (import_id, row_number, student_name, jshshir, error_message, raw_data)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (import_id, r.get("row_num", 0), r.get("full_name"), r.get("jshshir"), msg, json.dumps(r, ensure_ascii=False)))
                continue
                
            bed_id = bed['id']
            sport_id = sports_map.get(str(r.get("sport", "")).lower(), 1)
            priv_id = priv_map.get(str(r.get("privilege", "")).lower(), 7) # 7 is default "Oddiy"
            
            total_fee = float(r.get("total_fee", 2400000.0))
            paid_fee = float(r.get("paid_fee", 0.0))
            remaining_fee = max(0.0, total_fee - paid_fee)
            payment_status = "paid" if remaining_fee == 0 else ("partial" if paid_fee > 0 else "unpaid")
            
            stud_cursor = conn.execute("""
                INSERT INTO students (
                    full_name, jshshir, faculty, course, phone, email, gender,
                    sport_id, privilege_id, total_fee, paid_fee, remaining_fee,
                    payment_status, status, current_room_id, current_bed_id, check_in_date
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?, ?)
            """, (
                r.get("full_name"),
                r.get("jshshir"),
                r.get("faculty"),
                r.get("course"),
                r.get("phone", "+998901234567"),
                f"student_{r.get('jshshir')[-4:]}@med.uz",
                r.get("gender", "Erkak"),
                sport_id,
                priv_id,
                total_fee,
                paid_fee,
                remaining_fee,
                payment_status,
                room_id,
                bed_id,
                now_date
            ))
            student_id = stud_cursor.lastrowid
            
            # Occupy bed
            conn.execute("UPDATE room_beds SET status = 'occupied', current_student_id = ? WHERE id = ?", (student_id, bed_id))
            
            # Accommodation record
            conn.execute("""
                INSERT INTO student_accommodation (student_id, room_id, bed_id, check_in_date, status, reason, created_by)
                VALUES (?, ?, ?, ?, 'active', 'Ommaviy import orqali joylashtirildi', ?)
            """, (student_id, room_id, bed_id, now_date, current_user.get("full_name")))
            
            if paid_fee > 0:
                conn.execute("""
                    INSERT INTO payments (student_id, amount, payment_date, payment_method, notes, created_by)
                    VALUES (?, ?, ?, 'Bank o\'tkazmasi', 'Import orqali boshlang\'ich to\'lov', ?)
                """, (student_id, paid_fee, now_date, current_user.get("full_name")))
                
            success_count += 1
            
        status_val = "completed" if error_count == 0 else ("partial" if success_count > 0 else "failed")
        conn.execute("""
            UPDATE imports SET success_count = ?, error_count = ?, status = ? WHERE id = ?
        """, (success_count, error_count, status_val, import_id))
        
    log_audit("import_confirm", "import", import_id, None, {
        "file_name": file_name,
        "success_count": success_count,
        "error_count": error_count
    }, f"Ommaviy import tasdiqlandi: {success_count} ta talaba muvaffaqiyatli kiritildi, {error_count} ta xato",
    current_user.get("id"), current_user.get("full_name"))
    
    return {
        "success": True,
        "import_id": import_id,
        "success_count": success_count,
        "error_count": error_count,
        "status": status_val,
        "message": f"Import yakunlandi: {success_count} ta talaba joylashtirildi, {error_count} ta xatolik."
    }

def get_import_history(limit: int = 20, offset: int = 0) -> Dict[str, Any]:
    with get_db() as conn:
        imports = dicts_from_rows(conn.execute("""
            SELECT * FROM imports ORDER BY id DESC LIMIT ? OFFSET ?
        """, (limit, offset)).fetchall())
        for imp in imports:
            errors = dicts_from_rows(conn.execute("""
                SELECT * FROM import_errors WHERE import_id = ? ORDER BY row_number ASC LIMIT 50
            """, (imp['id'],)).fetchall())
            imp['errors'] = errors
        total = conn.execute("SELECT COUNT(*) as count FROM imports").fetchone()["count"]
        return {"total": total, "items": imports}

def generate_excel_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Talabalar_Shablon"
    
    headers = [
        "F.I.O.", "JShShIR", "Fakultet", "Kurs", "Xona", "Sport",
        "Imtiyoz", "Umumiy to'lov", "To'langan", "Telefon", "Jinsi"
    ]
    ws.append(headers)
    
    # Style headers
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # Sample rows
    samples = [
        ["Karimov Jasur Anvarovich", "30101990120034", "Davolash", 2, "101", "Futbol", "Oddiy / Imtiyozsiz", 2400000, 2400000, "+998901234567", "Erkak"],
        ["Ismoilova Nigora Rustam qizi", "40203010230045", "Pediatriya", 1, "102", "Shaxmat", "Yetimlik", 2400000, 1200000, "+998939876543", "Ayol"],
        ["Sobirov Eldor Baxtiyor o'g'li", "31205020340056", "Stomatologiya", 3, "201", "Voleybol", "Temir daftar", 2400000, 0, "+998945556677", "Erkak"]
    ]
    for row in samples:
        ws.append(row)
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def generate_csv_template() -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "F.I.O.", "JShShIR", "Fakultet", "Kurs", "Xona", "Sport",
        "Imtiyoz", "Umumiy to'lov", "To'langan", "Telefon", "Jinsi"
    ])
    writer.writerow(["Karimov Jasur Anvarovich", "30101990120034", "Davolash", "2", "101", "Futbol", "Oddiy / Imtiyozsiz", "2400000", "2400000", "+998901234567", "Erkak"])
    writer.writerow(["Ismoilova Nigora Rustam qizi", "40203010230045", "Pediatriya", "1", "102", "Shaxmat", "Yetimlik", "2400000", "1200000", "+998939876543", "Ayol"])
    return output.getvalue()
